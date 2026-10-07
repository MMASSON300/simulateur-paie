"""Extrait les données de référence du classeur Excel vers des fichiers JSON.

Ce script est exécuté une fois (hors ligne) pour produire les fichiers de seed
consommés par ``seed.py``. Il garantit la fidélité des données par rapport au
classeur source (aucune saisie manuelle).

Usage :
    python -m app.extract <chemin_vers_xlsm>
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

SEED_DIR = Path(__file__).parent / "seed_data"


def _num(v):
    """Retourne None si non numérique, sinon float."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f


def extract_grades(wb) -> list[dict]:
    bg = wb["Base_grades"]
    rows = []
    for r in range(2, bg.max_row + 1):
        filiere = bg.cell(r, 1).value
        cadre = bg.cell(r, 2).value
        grade = bg.cell(r, 4).value
        echelon = _num(bg.cell(r, 5).value)
        im = _num(bg.cell(r, 7).value)
        categorie = bg.cell(r, 8).value
        if grade is None or echelon is None or im is None:
            continue
        rows.append(
            {
                "filiere": filiere,
                "cadre": cadre,
                "grade": grade,
                "echelon": int(echelon),
                "im": im,
                "categorie": categorie,
            }
        )
    return rows


def extract_ifse(wb) -> list[dict]:
    bi = wb["Base_IFSE"]
    rows = []
    for r in range(2, bi.max_row + 1):
        code_regime = bi.cell(r, 1).value
        regime = bi.cell(r, 3).value
        grade_code = bi.cell(r, 4).value
        grade = bi.cell(r, 5).value
        montant = _num(bi.cell(r, 6).value)
        if grade is None or montant is None:
            continue
        rows.append(
            {
                "code_regime": code_regime,
                "regime": regime,
                "grade_code": grade_code,
                "grade": grade,
                "montant": montant,
            }
        )
    return rows


def extract_sft(wb) -> list[dict]:
    sft = wb["SFT"]
    rows = []
    # Nombre d'enfants en colonne D (lignes 4 à 13), élément fixe en G,
    # élément proportionnel en H.
    for r in range(4, 14):
        nb = _num(sft.cell(r, 4).value)
        if nb is None:
            continue
        fixe = _num(sft.cell(r, 7).value)
        prop = sft.cell(r, 8).value
        prop = None if prop == "-" else _num(prop)
        rows.append(
            {
                "enfants": int(nb),
                "element_fixe": fixe,
                "element_proportionnel": prop,
            }
        )
    return rows


def extract_constants(wb) -> dict:
    bg = wb["Base_grades"]
    sim = wb["Simulation"]

    # Temps partiel : colonne P (quotité) -> colonne Q (facteur)
    temps_partiel = []
    for r in range(3, 8):
        q = _num(bg.cell(r, 16).value)  # P
        f = _num(bg.cell(r, 17).value)  # Q
        if q is not None and f is not None:
            temps_partiel.append({"quotite": q, "facteur": f})

    # Transfert prime point : catégorie (A/B/C) -> montant annuel
    transfert = []
    for r in range(12, 15):
        cat = bg.cell(r, 16).value  # P
        val = _num(bg.cell(r, 17).value)  # Q
        if cat in ("A", "B", "C") and val is not None:
            transfert.append({"categorie": cat, "montant_annuel": val})

    return {
        "valeur_point": _num(sim.cell(22, 4).value),  # D22
        "plafond_ss": _num(bg.cell(1, 20).value),  # T1
        "smic": _num(bg.cell(3, 20).value),  # T3
        "prime_pfa_mensuel": _num(bg.cell(8, 20).value),  # T8 = 1460/12
        "prime_segur_mensuel": _num(bg.cell(12, 20).value),  # T12 = CTI Ségur
        "prime_pga_mensuel": _num(bg.cell(16, 20).value),  # T16 = Prime grand âge
        "temps_partiel": temps_partiel,
        "transfert_prime_point": transfert,
    }


def extract_dropdowns(grades: list[dict], ifse: list[dict]) -> dict:
    """Construit les listes hiérarchiques pour les listes déroulantes du frontend."""
    filieres = sorted({g["filiere"] for g in grades if g["filiere"]})
    cadre_by_filiere: dict[str, list[str]] = {}
    grade_by_cadre: dict[str, list[str]] = {}
    for g in grades:
        f, c, gr = g["filiere"], g["cadre"], g["grade"]
        if not f or not c:
            continue
        cadre_by_filiere.setdefault(f, [])
        if c not in cadre_by_filiere[f]:
            cadre_by_filiere[f].append(c)
        grade_by_cadre.setdefault(c, [])
        if gr not in grade_by_cadre[c]:
            grade_by_cadre[c].append(gr)

    regimes = sorted({i["regime"] for i in ifse if i["regime"]})
    return {
        "filieres": filieres,
        "cadres_par_filiere": cadre_by_filiere,
        "grades_par_cadre": grade_by_cadre,
        "regimes_ifse": regimes,
        "echelons": list(range(1, 14)),
        "statuts": ["Titulaire", "Non titulaire"],
    }


def main() -> None:
    if len(sys.argv) != 2:
        print("Usage: python -m app.extract <chemin.xlsm>")
        sys.exit(1)

    path = Path(sys.argv[1])
    wb = openpyxl.load_workbook(path, data_only=True)

    grades = extract_grades(wb)
    ifse = extract_ifse(wb)
    sft = extract_sft(wb)
    constants = extract_constants(wb)
    dropdowns = extract_dropdowns(grades, ifse)

    SEED_DIR.mkdir(parents=True, exist_ok=True)

    def dump(name: str, data):
        (SEED_DIR / name).write_text(
            json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print(f"{name}: {len(data) if isinstance(data, list) else 'ok'} entrées")

    dump("grades.json", grades)
    dump("ifse.json", ifse)
    dump("sft.json", sft)
    dump("constants.json", constants)
    dump("dropdowns.json", dropdowns)


if __name__ == "__main__":
    main()
